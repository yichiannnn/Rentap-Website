#!/usr/bin/env python3
"""Build the fixtures seed from the organisers' schedule workbook.

Usage:
    python3 scripts/build-fixtures.py [path/to/schedule.xlsx] [-o data/fixtures-seed.json]

Reads the football/rugby, volleyball, badminton and table-tennis sheets, checks them
(round robins complete, no venue or player double bookings, expected match counts,
and an exact comparison with the workbook's own 'Player Timetable' sheet), then writes
one JSON document that scripts/seed-fixtures.mjs loads into the database through
/api/scores-admin.

Also writes data/athletes-seed.json — every entrant's gender for the admin-only Best
Athlete page — from the Name and Gender columns of the 'All Participant' roster, with
the gender implied by the competition (men's singles, women's doubles, "100m Women" …)
for the few entrants the roster does not list.

The output goes under data/ (gitignored) because it carries players' full names. The
script never copies the workbook and never reads the PIC/phone cells; from
'All Participant' it reads only columns A and B.
"""
import argparse
import collections
import datetime
import glob
import itertools
import json
import os
import re
import sys
import unicodedata

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DAY_DATES = {1: '2026-09-26', 2: '2026-09-27'}
TZ = '+02:00'  # Europe/Berlin is on CEST on both days

# distinct badge colours, assigned in sheet order per sport
PALETTE = ['#A8200D', '#C8922A', '#1E4D8C', '#0F8A4D', '#6A3D9A',
           '#D2691E', '#2A9D8F', '#B5179E', '#3A86FF', '#8D5524']
CATEGORY_COLOUR = {'MS': '#C8922A', 'MD': '#A8200D', 'XD': '#EDE8DE',
                   'WD': '#0F8A4D', 'WS': '#7A8572', 'OD': '#A8200D'}
CATEGORY_NAME = {'MS': "Men's Singles", 'MD': "Men's Doubles", 'XD': 'Mixed Doubles',
                 'WD': "Women's Doubles", 'WS': "Women's Singles", 'OD': 'Open Doubles'}

STAGE_RANK = {'group': 0, 'quarter': 1, 'semi': 2, 'third': 3, 'final': 4}
SPORT_NAME = {'football': 'Football', 'touch-rugby': 'Rugby Touch', 'volleyball': 'Volleyball',
              'badminton': 'Badminton', 'table-tennis': 'Table Tennis'}

EXPECTED = {  # slug: (teams/entries, group matches, knockout matches)
    'football': (10, 20, 4), 'touch-rugby': (3, 6, 0), 'volleyball': (6, 15, 0),
    'badminton-ms': (13, 22, 4), 'badminton-md': (15, 21, 8), 'badminton-xd': (7, 9, 4),
    'badminton-wd': (5, 10, 0), 'badminton-ws': (5, 10, 0),
    'table-tennis-ms': (8, 12, 4), 'table-tennis-ws': (4, 6, 2), 'table-tennis-od': (11, 25, 4),
}
EXPECTED_ORACLE = {  # (sport name, certainty): rows in the Player Timetable sheet
    ('Football', 'fixed'): 492, ('Football', 'if qualified'): 492,
    ('Rugby Touch', 'fixed'): 112, ('Volleyball', 'fixed'): 235,
    ('Badminton', 'fixed'): 224, ('Badminton', 'if qualified'): 348,
    ('Table Tennis', 'fixed'): 136, ('Table Tennis', 'if qualified'): 128,
}


class SheetError(Exception):
    pass


def check(cond, msg):
    if not cond:
        raise SheetError(msg)


def t2m(v):
    """Excel time cell -> minutes since midnight, else None."""
    if isinstance(v, (datetime.time, datetime.datetime)):
        return v.hour * 60 + v.minute
    return None


def hm(m):
    return f'{m // 60:02d}:{m % 60:02d}'


def s(v):
    """Cell value as a stripped string (or the value itself when not a string)."""
    return v.strip() if isinstance(v, str) else v


def family_of(slug):
    return slug.rsplit('-', 1)[0] if slug.startswith(('badminton-', 'table-tennis-')) else slug


def match(slug, day, time, duration, stage, number, group, a, b, venue, referee=None, half=None,
          placeholder_a=None, placeholder_b=None):
    return dict(slug=slug, day=day, time=time, duration=duration, stage=stage, number=number,
                group=group, a=a, b=b, venue=venue, referee=referee, half=half,
                placeholder_a=placeholder_a, placeholder_b=placeholder_b)


# ───────────────────────── Football & Rugby ─────────────────────────
FIELD_RE = re.compile(r'^(Football|Rugby)(?:\s+Group\s+([A-Z])|\s+Semi Final\s+(\d)|\s+3rd Place|\s+Final)?\s*\((\d+) min\)$')


def parse_field_sheet(wb):
    ws = wb['Football and Rugby Touch Sched']
    for c, want in [('A11', 'Time'), ('B11', 'Match type'), ('C11', 'Team Name 1'), ('D11', 'Team Name 2'),
                    ('E11', 'Field Number'), ('F11', 'Referee'), ('K11', 'Time'), ('L11', 'Match type'),
                    ('M11', 'Team Name 1'), ('N11', 'Team Name 2'), ('O11', 'Field Number'), ('P11', 'Referee'),
                    ('A37', 'Football Group A'), ('H37', 'Football Group B')]:
        check(s(ws[c].value) == want, f'football sheet: expected {c}={want!r}, got {ws[c].value!r}')
    check(str(ws['A10'].value).startswith('Day 1') and str(ws['K10'].value).startswith('Day 2'),
          'football sheet: day headers moved')
    check(str(ws['O37'].value).startswith('Rugby Touch'), 'football sheet: rugby table moved')
    check(str(ws['A46'].value).startswith('Squads'), 'football sheet: squads block moved')

    groups = {'A': [s(ws.cell(38, c).value) for c in range(2, 7)],
              'B': [s(ws.cell(38, c).value) for c in range(9, 14)]}
    rugby_teams = [s(ws.cell(38, c).value) for c in range(16, 19)]
    check(all(groups['A']) and all(groups['B']) and all(rugby_teams), 'football sheet: empty team in a group table')
    group_of = {t: g for g, ts in groups.items() for t in ts}

    squads = {}
    for c in range(1, 14):
        name = s(ws.cell(47, c).value)
        check(name, f'football sheet: empty squad header in column {c}')
        squads[name] = [s(ws.cell(r, c).value) for r in range(48, ws.max_row + 1) if s(ws.cell(r, c).value)]
    check(set(squads) == set(group_of) | set(rugby_teams),
          f'football sheet: squad headers {sorted(squads)} differ from the team tables')

    matches = []
    for day, cols in ((1, 'ABCDEF'), (2, 'KLMNOP')):
        for r in range(12, 45):
            t = t2m(ws[f'{cols[0]}{r}'].value)
            if t is None:
                if str(ws[f'{cols[0]}{r}'].value).startswith('Games end'):
                    break
                continue  # BREAK rows and blanks
            mt = s(ws[f'{cols[1]}{r}'].value)
            if not mt or mt.startswith('Games end'):
                break
            m = FIELD_RE.match(mt)
            check(m, f'football sheet: unrecognised match type {mt!r} at {cols[1]}{r}')
            kind, grp, sfn, dur = m.group(1), m.group(2), m.group(3), int(m.group(4))
            a, b = s(ws[f'{cols[2]}{r}'].value), s(ws[f'{cols[3]}{r}'].value)
            venue, ref = s(ws[f'{cols[4]}{r}'].value), s(ws[f'{cols[5]}{r}'].value)
            check(venue, f'football sheet: missing field at {cols[4]}{r}')
            where = f'{cols[1]}{r} ({mt})'
            if kind == 'Rugby':
                check(a in rugby_teams and b in rugby_teams, f'football sheet: unknown rugby team at {where}')
                matches.append(match('touch-rugby', day, t, dur, 'group', None, None, a, b, venue, ref))
            elif grp:
                check(group_of.get(a) == grp and group_of.get(b) == grp,
                      f'football sheet: {a!r} v {b!r} are not both in group {grp} at {where}')
                matches.append(match('football', day, t, dur, 'group', None, grp, a, b, venue, ref, half=10))
            elif sfn:
                check(a and b and a not in group_of and b not in group_of,
                      f'football sheet: semi final at {where} should carry placeholder names')
                matches.append(match('football', day, t, dur, 'semi', int(sfn), None, None, None, venue, ref,
                                     half=10, placeholder_a=a, placeholder_b=b))
            elif '3rd Place' in mt:
                check(not a and not b, f'football sheet: unexpected teams at {where}')
                matches.append(match('football', day, t, dur, 'third', 1, None, None, None, venue, ref,
                                     half=10, placeholder_a='Loser SF1', placeholder_b='Loser SF2'))
            else:
                check(not a and not b, f'football sheet: unexpected teams at {where}')
                # the sheet only says "60 minute slot"; 15-minute halves is an assumption
                matches.append(match('football', day, t, dur, 'final', 1, None, None, None, venue, ref,
                                     half=15, placeholder_a='Winner SF1', placeholder_b='Winner SF2'))

    teams = {
        'football': [dict(name=t, code=None, group=group_of[t], color=PALETTE[i % len(PALETTE)], players=squads[t])
                     for i, t in enumerate(groups['A'] + groups['B'])],
        'touch-rugby': [dict(name=t, code=None, group=None, color=PALETTE[i], players=squads[t])
                        for i, t in enumerate(rugby_teams)],
    }
    return teams, matches


# ───────────────────────── Volleyball ─────────────────────────
def parse_volleyball(wb):
    ws = wb['Volleyball Schedule']
    for c, want in [('A12', 'Time'), ('B12', 'Match type'), ('C12', 'Team Name 1'), ('D12', 'Team Name 2'),
                    ('E12', 'Court'), ('I12', 'Time'), ('J12', 'Match type'), ('K12', 'Team Name 1'),
                    ('L12', 'Team Name 2'), ('M12', 'Court'), ('A25', 'League table'), ('A35', 'Squads')]:
        check(s(ws[c].value) == want, f'volleyball sheet: expected {c}={want!r}, got {ws[c].value!r}')
    check(str(ws['A11'].value).startswith('Day 1') and str(ws['I11'].value).startswith('Day 2'),
          'volleyball sheet: day headers moved')
    check(str(ws['F12'].value).startswith('Referee') and str(ws['N12'].value).startswith('Referee'),
          'volleyball sheet: referee columns moved')

    league = [s(ws.cell(26, c).value) for c in range(2, 8)]
    check(all(league), 'volleyball sheet: empty team in the league table')
    squads = {}
    for c in range(1, 7):
        name = s(ws.cell(36, c).value)
        check(name, f'volleyball sheet: empty squad header in column {c}')
        squads[name] = [s(ws.cell(r, c).value) for r in range(37, ws.max_row + 1) if s(ws.cell(r, c).value)]
    check(set(squads) == set(league), f'volleyball sheet: squad headers {sorted(squads)} differ from the league table')

    matches = []
    for day, cols in ((1, 'ABCDEF'), (2, 'IJKLMN')):
        for r in range(13, 35):
            t = t2m(ws[f'{cols[0]}{r}'].value)
            if t is None:
                break
            mt, a, b = s(ws[f'{cols[1]}{r}'].value), s(ws[f'{cols[2]}{r}'].value), s(ws[f'{cols[3]}{r}'].value)
            venue, ref = s(ws[f'{cols[4]}{r}'].value), s(ws[f'{cols[5]}{r}'].value)
            check(mt == 'League', f'volleyball sheet: unexpected match type {mt!r} at {cols[1]}{r}')
            check(a in league and b in league, f'volleyball sheet: unknown team at {cols[2]}{r}')
            check(venue, f'volleyball sheet: missing court at {cols[4]}{r}')
            matches.append(match('volleyball', day, t, 60, 'group', None, None, a, b, venue, ref))

    teams = {'volleyball': [dict(name=t, code=None, group=None, color=PALETTE[i], players=squads[t])
                            for i, t in enumerate(league)]}
    return teams, matches


# ───────────────────────── Badminton & Table tennis ─────────────────────────
CODE_RE = re.compile(r'^(MS|MD|XD|WD|WS|OD)(\d+)$')
KO_RE = re.compile(r'^(?P<cat>MS|MD|XD|WD|WS|OD)\s+(?P<stage>QF|Semi Final|Final|3rd Place)\s*(?P<n>\d)?'
                   r'\s*(?:[:(]\s*(?P<a>.+?)\s+vs\s+(?P<b>.+?)\)?)?\s*$')
KO_STAGE = {'QF': 'quarter', 'Semi Final': 'semi', 'Final': 'final', '3rd Place': 'third'}
COURT_PAIRS = [(2, 3), (6, 7), (10, 11), (14, 15), (18, 19), (22, 23)]  # (left, right) columns per court


def parse_entries(ws, family, rows, triples):
    """triples: (category, group column, code column, name column). Returns {code: entry}."""
    entries = {}
    for cat, gc, cc, nc in triples:
        for r in rows:
            g, code, name = s(ws[f'{gc}{r}'].value), s(ws[f'{cc}{r}'].value), s(ws[f'{nc}{r}'].value)
            if not code and not name and not g:
                continue
            check(g and code and name, f'{family}: incomplete {cat} entry in row {r}')
            m = CODE_RE.match(code)
            check(m and m.group(1) == cat, f'{family}: code {code!r} found in the {cat} column (row {r})')
            check(code not in entries, f'{family}: duplicate code {code}')
            players = [p.strip() for p in name.split(' / ')]
            check(1 <= len(players) <= 2 and all(players), f'{family}: cannot read the names in {name!r}')
            entries[code] = dict(cat=cat, code=code, group=None if g == 'RR' else g, name=name, players=players)
    for cat in {e['cat'] for e in entries.values()}:
        sizes = {len(e['players']) for e in entries.values() if e['cat'] == cat}
        check(len(sizes) == 1, f'{family}: {cat} mixes singles and doubles entries')
    return entries


def parse_grid(ws, family, entries, day, header_row, first_row, hall_for, venue_word, matches):
    # the "Time" header sits at header_row in the v6 workbook; tolerate a row or two of drift
    found = next((r for r in (header_row, header_row - 1, header_row + 1, header_row - 2, header_row + 2)
                  if s(ws[f'A{r}'].value) == 'Time'), None)
    check(found is not None, f'{family} day {day}: header row moved (expected "Time" near A{header_row})')
    first_row += found - header_row
    header_row = found
    for i, (lc, _) in enumerate(COURT_PAIRS):
        hdr = s(ws.cell(header_row, lc).value)
        if hdr is not None:
            check(hdr.startswith(f'{venue_word} {i + 1}'),
                  f'{family} day {day}: expected {venue_word} {i + 1} header, got {hdr!r}')
    rows = []
    r = first_row
    while True:
        t = t2m(ws[f'A{r}'].value)
        check(t is not None, f'{family} day {day}: expected a time at A{r}, got {ws[f"A{r}"].value!r}')
        rows.append((r, t))
        label = s(ws.cell(r, 2).value)
        if isinstance(label, str) and label.lower().endswith('ends'):
            break
        r += 1
        check(r < first_row + 40, f'{family} day {day}: no "ends" row found')
    for idx, (r, t) in enumerate(rows[:-1]):
        dur = rows[idx + 1][1] - t
        check(dur > 0, f'{family} day {day}: times do not increase at row {r}')
        for ci, (lc, rc) in enumerate(COURT_PAIRS):
            left, right = s(ws.cell(r, lc).value), s(ws.cell(r, rc).value)
            if left is None and right is None:
                continue
            if left in ('Break', 'Basketball'):
                check(right is None, f'{family} day {day} row {r}: unexpected value next to {left!r}')
                continue
            venue = f'{hall_for(day, ci + 1)} · {venue_word} {ci + 1}'
            where = f'{family} day {day} {hm(t)} {venue}'
            if CODE_RE.match(str(left)):
                check(right and CODE_RE.match(str(right)), f'{where}: {left!r} has no opponent')
                ea, eb = entries.get(left), entries.get(right)
                check(ea and eb, f'{where}: unknown code {left!r} / {right!r}')
                check(ea['cat'] == eb['cat'] and ea['group'] == eb['group'],
                      f'{where}: {left} and {right} are not in the same group')
                matches.append(match(f'{family}-{ea["cat"].lower()}', day, t, dur, 'group', None, ea['group'],
                                     left, right, venue))
            else:
                check(right is None, f'{where}: unexpected second cell {right!r} next to {left!r}')
                m = KO_RE.match(str(left))
                check(m, f'{where}: unrecognised cell {left!r}')
                stage = KO_STAGE[m.group('stage')]
                pa, pb = m.group('a'), m.group('b')
                if not pa:
                    pa, pb = ('Winner SF1', 'Winner SF2') if stage == 'final' else ('Loser SF1', 'Loser SF2')
                matches.append(match(f'{family}-{m.group("cat").lower()}', day, t, dur, stage,
                                     int(m.group('n')) if m.group('n') else 1, None, None, None, venue,
                                     placeholder_a=pa, placeholder_b=pb))


def racket_teams(family, entries):
    teams = collections.defaultdict(list)
    for e in entries.values():
        teams[f'{family}-{e["cat"].lower()}'].append(
            dict(name=e['name'], code=e['code'], group=e['group'], color=CATEGORY_COLOUR[e['cat']], players=e['players']))
    return dict(teams)


def parse_badminton(wb):
    ws = wb['Badminton Schedule']
    for c, want in [('A9', "Men's Singles"), ('E9', "Men's Doubles"), ('I9', 'Mixed Doubles'),
                    ('M9', "Women's Doubles"), ('Q9', "Women's Singles"), ('A10', 'Group'), ('B10', 'Code'),
                    ('C10', 'Name')]:
        check(s(ws[c].value) == want, f'badminton sheet: expected {c}={want!r}, got {ws[c].value!r}')
    entries = parse_entries(ws, 'badminton', range(11, 26),
                            [('MS', 'A', 'B', 'C'), ('MD', 'E', 'F', 'G'), ('XD', 'I', 'J', 'K'),
                             ('WD', 'M', 'N', 'O'), ('WS', 'Q', 'R', 'S')])
    matches = []
    hall = lambda day, court: 'Hall 1' if court <= 3 else 'Hall 3'
    parse_grid(ws, 'badminton', entries, 1, 37, 38, hall, 'Court', matches)
    parse_grid(ws, 'badminton', entries, 2, 60, 61, hall, 'Court', matches)
    return racket_teams('badminton', entries), matches


def parse_table_tennis(wb):
    ws = wb['Pingpong Schedule']
    for c, want in [('A8', "Men's Singles"), ('E8', "Women's Singles"), ('I8', 'Open Doubles'),
                    ('A9', 'Group'), ('B9', 'Code'), ('C9', 'Name')]:
        check(s(ws[c].value) == want, f'table tennis sheet: expected {c}={want!r}, got {ws[c].value!r}')
    entries = parse_entries(ws, 'table-tennis', range(10, 21),
                            [('MS', 'A', 'B', 'C'), ('WS', 'E', 'F', 'G'), ('OD', 'I', 'J', 'K')])
    matches = []
    hall = lambda day, table: 'Hall 4' if day == 1 else 'Hall 3'
    parse_grid(ws, 'table-tennis', entries, 1, 31, 32, hall, 'Table', matches)
    parse_grid(ws, 'table-tennis', entries, 2, 49, 50, hall, 'Table', matches)
    return racket_teams('table-tennis', entries), matches


# ───────────────────────── Checks ─────────────────────────
def side_key(team):
    return team['code'] or team['name']


def verify(sports):
    problems = []

    # every side resolves; group matches stay inside their group; counts
    for slug, sp in sports.items():
        keys = {side_key(t): t for t in sp['teams']}
        check(len(keys) == len(sp['teams']), f'{slug}: duplicate team/entry names')
        for m in sp['matches']:
            if m['a'] is not None:
                check(m['a'] in keys and m['b'] in keys, f'{slug}: unknown side in {m}')
                check(m['a'] != m['b'], f'{slug}: a team plays itself: {m}')
                if m['stage'] == 'group':
                    check(keys[m['a']]['group'] == m['group'] == keys[m['b']]['group'],
                          f'{slug}: group mismatch in {m}')
            else:
                check(m['stage'] != 'group' and m['placeholder_a'] and m['placeholder_b'],
                      f'{slug}: unresolved teams in {m}')
        groups = sum(1 for m in sp['matches'] if m['stage'] == 'group')
        kos = len(sp['matches']) - groups
        want = EXPECTED[slug]
        check((len(sp['teams']), groups, kos) == want,
              f'{slug}: got {len(sp["teams"])} teams / {groups} group / {kos} knockout matches, expected {want}')

        # round robins: every pair exactly once (rugby: twice)
        need = 2 if slug == 'touch-rugby' else 1
        by_group = collections.defaultdict(list)
        for t in sp['teams']:
            by_group[t['group']].append(side_key(t))
        pair_count = collections.Counter(frozenset((m['a'], m['b'])) for m in sp['matches'] if m['stage'] == 'group')
        for g, members in by_group.items():
            for pair in itertools.combinations(members, 2):
                n = pair_count.get(frozenset(pair), 0)
                if n != need:
                    problems.append(f'{slug} group {g or "RR"}: {pair[0]} v {pair[1]} scheduled {n}× (need {need})')
        check(sum(pair_count.values()) == sum(len(v) * (len(v) - 1) // 2 * need for v in by_group.values()),
              f'{slug}: stray group matches outside the round robin')

    # venues: one match at a time (all sports share the fields)
    by_venue = collections.defaultdict(list)
    for slug, sp in sports.items():
        for m in sp['matches']:
            by_venue[(m['day'], m['venue'])].append((m['time'], m['time'] + m['duration'], slug, m))
    for (day, venue), lst in by_venue.items():
        lst.sort()
        for (s1, e1, slug1, _), (s2, e2, slug2, _) in zip(lst, lst[1:]):
            if s2 < e1:
                problems.append(f'day {day} {venue}: {slug1} {hm(s1)}–{hm(e1)} overlaps {slug2} {hm(s2)}–{hm(e2)}')

    # people: no two fixed matches at once, across all sports
    busy = collections.defaultdict(list)
    for slug, sp in sports.items():
        members = {side_key(t): t['players'] for t in sp['teams']}
        for m in sp['matches']:
            if m['a'] is None:
                continue
            for side in (m['a'], m['b']):
                for p in members[side]:
                    busy[(p, m['day'])].append((m['time'], m['time'] + m['duration'], slug, side))
    for (p, day), lst in busy.items():
        lst.sort()
        for (s1, e1, sl1, t1), (s2, e2, sl2, t2) in zip(lst, lst[1:]):
            if s2 < e1:
                problems.append(f'{p} day {day}: {sl1} {t1} {hm(s1)}–{hm(e1)} overlaps {sl2} {t2} {hm(s2)}–{hm(e2)}')

    # names never carry numbers (guards against phone numbers or codes leaking into names)
    for slug, sp in sports.items():
        for t in sp['teams']:
            for name in [t['name']] + t['players']:
                check(not re.search(r'\d{3,}', name), f'{slug}: suspicious name {name!r}')

    check(not problems, 'schedule problems:\n  ' + '\n  '.join(problems))


def oracle(wb, sports):
    """Compare our per-person view with the workbook's own 'Player Timetable' sheet."""
    ws = wb['Player Timetable']
    check([s(c.value) for c in ws[2]] == ['Player', 'Registrations', 'Day', 'Start', 'End', 'Sport', 'Event',
                                          'Venue', 'Certainty'], 'Player Timetable: header row moved')
    expected, current = set(), None
    for r in ws.iter_rows(min_row=3, values_only=True):
        if all(v is None for v in r):
            continue
        if r[0]:
            current = s(r[0])
        if r[5] not in SPORT_NAME.values():
            continue  # track, tug of war, frisbee, basketball are not seeded
        expected.add((current, r[5], int(str(r[2]).split()[-1]), t2m(r[3]), t2m(r[4]), s(r[7]), s(r[8])))

    derived = set()
    for slug, sp in sports.items():
        sport = SPORT_NAME[family_of(slug)]
        members = {side_key(t): t['players'] for t in sp['teams']}
        everyone = {p for ps in members.values() for p in ps}
        for m in sp['matches']:
            venue = m['venue'].replace(' · ', ' ')
            end = m['time'] + m['duration']
            if m['a'] is not None:
                for side in (m['a'], m['b']):
                    for p in members[side]:
                        derived.add((p, sport, m['day'], m['time'], end, venue, 'fixed'))
            else:  # knockout: every entrant of this competition may still be in it
                for p in everyone:
                    derived.add((p, sport, m['day'], m['time'], end, venue, 'if qualified'))

    fmt = lambda row: f'{row[0]} | {row[1]} | day {row[2]} {hm(row[3])}–{hm(row[4])} | {row[5]} | {row[6]}'
    missing, extra = sorted(expected - derived), sorted(derived - expected)
    if missing or extra:
        lines = [f'  sheet has, we lack: {fmt(x)}' for x in missing[:20]] + \
                [f'  we have, sheet lacks: {fmt(x)}' for x in extra[:20]]
        check(False, f'Player Timetable mismatch ({len(missing)} missing, {len(extra)} extra):\n' + '\n'.join(lines))
    counts = collections.Counter((row[1], row[6]) for row in derived)
    check(dict(counts) == EXPECTED_ORACLE, f'Player Timetable row counts changed: {dict(counts)}')
    return counts


# ───────────────────────── Athletes (gender) ─────────────────────────
def norm_name(v):
    """Same identity rule as lib/names.js normName(): no accents, lower case, single spaces."""
    text = unicodedata.normalize('NFD', str(v or ''))
    text = ''.join(ch for ch in text if unicodedata.category(ch) != 'Mn')
    return re.sub(r'\s+', ' ', text.lower()).strip()


def build_athletes(wb, sports):
    """[{name, gender}] for every entrant: the roster's own M/F, or the gender the
    competition implies for entrants the roster does not list. Returns
    (athletes, inferred, unknown)."""
    ws = wb['All Participant']
    check(s(ws['A1'].value) == 'Name' and s(ws['B1'].value) == 'Gender',
          f'All Participant: expected Name / Gender in A1 / B1, got {ws["A1"].value!r} / {ws["B1"].value!r}')
    roster = {}
    for r in range(2, ws.max_row + 1):
        name, gender = s(ws.cell(r, 1).value), s(ws.cell(r, 2).value)
        if not name:
            continue
        check(gender in ('M', 'F'), f'All Participant row {r}: gender {gender!r} is not M or F')
        roster[norm_name(name)] = dict(name=name, gender=gender)

    # entrants with the gender their competition implies (None when it implies nothing)
    implied = {}
    for slug, sp in sports.items():
        gender = ('M' if slug.endswith(('-ms', '-md')) or slug == 'football'
                  else 'F' if slug.endswith(('-ws', '-wd')) else None)
        for t in sp['teams']:
            for p in t['players']:
                implied.setdefault(norm_name(p), dict(name=p, gender=gender))
    tws = wb['Track Schedule']   # runner lists per event, headed "100m Men", "4x100m Women (…)" …
    check(str(tws['A23'].value).startswith('Player'), 'Track Schedule: the player list moved (expected it at A23)')
    for c in range(1, tws.max_column + 1):
        hdr = s(tws.cell(24, c).value)
        if not hdr:
            continue
        gender = 'F' if 'women' in hdr.lower() else 'M' if 'men' in hdr.lower() else None
        for r in range(25, tws.max_row + 1):
            name = s(tws.cell(r, c).value)
            if name:
                implied.setdefault(norm_name(name), dict(name=name, gender=gender))

    athletes = [dict(name=v['name'], gender=v['gender']) for v in roster.values()]
    inferred, unknown = [], []
    for key, v in implied.items():
        if key in roster:
            continue
        if v['gender']:
            athletes.append(dict(name=v['name'], gender=v['gender']))
            inferred.append(f"{v['name']} → {v['gender']}")
        else:
            unknown.append(v['name'])
    return athletes, inferred, unknown


# ───────────────────────── Main ─────────────────────────
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('xlsx', nargs='?', help='schedule workbook (default: the only .xlsx under data/)')
    ap.add_argument('-o', '--out', default=os.path.join(ROOT, 'data', 'fixtures-seed.json'))
    ap.add_argument('-a', '--athletes-out', default=os.path.join(ROOT, 'data', 'athletes-seed.json'))
    args = ap.parse_args()

    path = args.xlsx
    if not path:
        found = glob.glob(os.path.join(ROOT, 'data', '*.xlsx'))
        if len(found) != 1:
            sys.exit(f'expected exactly one .xlsx under data/, found {len(found)}; pass the path explicitly')
        path = found[0]

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        sports = {}
        all_matches = []
        for teams, matches in (parse_field_sheet(wb), parse_volleyball(wb), parse_badminton(wb), parse_table_tennis(wb)):
            for slug, ts in teams.items():
                sports[slug] = dict(family=family_of(slug), teams=ts, matches=[])
            all_matches += matches
        for m in all_matches:
            check(m['slug'] in sports, f'match for unknown sport {m["slug"]}')
            sports[m['slug']]['matches'].append(m)
        for slug, sp in sports.items():
            # creation order = bracket order: group matches chronologically, then knockouts by round and number
            sp['matches'].sort(key=lambda m: (STAGE_RANK[m['stage']] > 0, STAGE_RANK[m['stage']], m['number'] or 0,
                                              m['day'], m['time'], m['venue']))
            for m in sp['matches']:
                m['time'] = hm(m['time'])
                del m['slug']
        verify_input = {slug: dict(sp, matches=[dict(m, time=int(m['time'][:2]) * 60 + int(m['time'][3:]))
                                                  for m in sp['matches']]) for slug, sp in sports.items()}
        verify(verify_input)
        counts = oracle(wb, verify_input)
        athletes, inferred, unknown = build_athletes(wb, sports)
    except SheetError as e:
        sys.exit(f'build-fixtures: {e}')

    out = dict(
        meta=dict(source=os.path.basename(path), generated=datetime.datetime.now().isoformat(timespec='seconds'),
                  tz=TZ, days=DAY_DATES,
                  categories={slug: CATEGORY_NAME[slug.rsplit('-', 1)[1].upper()]
                              for slug in sports if family_of(slug) != slug},
                  counts={slug: dict(teams=len(sp['teams']), matches=len(sp['matches'])) for slug, sp in sports.items()}),
        sports=sports,
    )
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    with open(args.out, 'w', encoding='utf-8') as fh:
        json.dump(out, fh, ensure_ascii=False, indent=1)

    print(f'{"sport":18} {"teams":>5} {"group":>5} {"KO":>3}   day 1 / day 2')
    for slug, sp in sports.items():
        groups = sum(1 for m in sp['matches'] if m['stage'] == 'group')
        d1 = sum(1 for m in sp['matches'] if m['day'] == 1)
        print(f'{slug:18} {len(sp["teams"]):5} {groups:5} {len(sp["matches"]) - groups:3}   {d1:3} / {len(sp["matches"]) - d1}')
    print('Player Timetable check: ' + ', '.join(f'{k[0]} {k[1]} {v}' for k, v in sorted(counts.items())))
    print(f'wrote {os.path.relpath(args.out, ROOT)}')

    with open(args.athletes_out, 'w', encoding='utf-8') as fh:
        json.dump(athletes, fh, ensure_ascii=False, indent=1)
    by_gender = collections.Counter(a['gender'] for a in athletes)
    print(f'athletes: {len(athletes)} ({by_gender["M"]} M / {by_gender["F"]} F), '
          f'{len(inferred)} not on the roster with gender implied by their competition'
          + (': ' + '; '.join(inferred) if inferred else ''))
    if unknown:
        print(f'  {len(unknown)} entrant(s) without a gender — set it on awards-admin.html: ' + '; '.join(unknown))
    print(f'wrote {os.path.relpath(args.athletes_out, ROOT)}')


if __name__ == '__main__':
    main()
